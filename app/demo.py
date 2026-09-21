"""演示合同:用 python-docx 生成一份带四类问题的采购合同,走 MinerU 解析入库。
四类问题:付款账期 90 天、违约金 30%、管辖约定深圳、全篇无保密条款。
答辩演示前跑一次,之后对话/审查全部走库,不再现场解析。
"""
from pathlib import Path

from docx import Document

from . import db
from .config import DATA_DIR
from .parse_pdf import parse_and_store

DEMO_CLAUSES = [
    ("第一条", "合同主体", "甲方:北京华信科技有限公司,乙方:上海启帆贸易有限公司。双方本着平等自愿原则,就设备采购事宜订立本合同。"),
    ("第二条", "合同金额", "本合同总金额为人民币 500,000 元(含税)。"),
    ("第三条", "付款方式", "甲方应于本合同生效后 90 日内向乙方支付全部货款。付款方式为银行转账,乙方应在收款前开具增值税专用发票。"),
    ("第四条", "交付", "乙方应于本合同生效后 30 日内将全部设备交付至甲方指定仓库,交付前风险由乙方承担。"),
    ("第五条", "质量与验收", "设备质量标准以双方确认的技术规格书为准。甲方应在收到设备后 10 个工作日内完成验收,质保期为验收合格后 12 个月。"),
    ("第六条", "违约责任", "任何一方违约,应向守约方支付合同总金额 30% 的违约金。违约金不足以弥补损失的,守约方有权要求按实际损失赔偿,包括律师费、诉讼费。"),
    ("第七条", "争议解决", "因本合同引起的争议,双方应先行协商;协商不成的,提交深圳市南山区人民法院诉讼解决。本合同适用中华人民共和国法律。"),
    ("第八条", "其他", "本合同一式两份,双方各执一份,自双方签字盖章之日起生效。"),
]


def build_demo_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("设备采购合同", level=1)
    for num, title, body in DEMO_CLAUSES:
        doc.add_heading(f"{num} {title}", level=2)
        doc.add_paragraph(body)
    doc.save(str(path))


def prepare_demo() -> int:
    """生成演示合同 docx → MinerU 解析 → 落库。返回 contract_id。"""
    docx_path = DATA_DIR / "demo_contract.docx"
    build_demo_docx(docx_path)
    file_id = db.insert_file(docx_path.name, str(docx_path), docx_path.stat().st_size)
    contract_id = parse_and_store(file_id, str(docx_path), docx_path.name)
    print(f"[demo] 演示合同已入库:contract_id={contract_id}, 条款 {db.get_contract(contract_id)['clause_count']} 条")
    return contract_id


if __name__ == "__main__":
    prepare_demo()
